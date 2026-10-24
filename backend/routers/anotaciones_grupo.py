from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from datetime import datetime
import io

from database import SessionLocal
import models
import schemas

router = APIRouter(prefix="/anotaciones-grupo", tags=["anotaciones_grupo"])


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


# ─────────────────────────────────────────
# GET /anotaciones-grupo/{grupo_id}
# Devuelve la anotación completa del grupo (crea una vacía si no existe)
# ─────────────────────────────────────────
@router.get("/{grupo_id}", response_model=schemas.AnotacionGrupoOut)
def get_anotacion(grupo_id: int, db: Session = Depends(get_db)):
    grupo = db.query(models.Grupo).filter(models.Grupo.id == grupo_id).first()
    if not grupo:
        raise HTTPException(status_code=404, detail="Grupo no encontrado")

    anotacion = db.query(models.AnotacionGrupo).filter(
        models.AnotacionGrupo.grupo_id == grupo_id
    ).first()

    if not anotacion:
        anotacion = models.AnotacionGrupo(grupo_id=grupo_id)
        db.add(anotacion)
        db.commit()
        db.refresh(anotacion)

    # Enriquecer con datos del grupo/empresa
    out = schemas.AnotacionGrupoOut.model_validate(anotacion)
    out.grupo_nombre = grupo.nombre
    out.empresa_nombre = grupo.empresa.nombre if grupo.empresa else None
    return out


# ─────────────────────────────────────────
# PUT /anotaciones-grupo/{grupo_id}
# Actualiza composición + notas generales
# ─────────────────────────────────────────
@router.put("/{grupo_id}", response_model=schemas.AnotacionGrupoOut)
def update_anotacion(grupo_id: int, body: schemas.AnotacionGrupoUpdate, db: Session = Depends(get_db)):
    grupo = db.query(models.Grupo).filter(models.Grupo.id == grupo_id).first()
    if not grupo:
        raise HTTPException(status_code=404, detail="Grupo no encontrado")

    anotacion = db.query(models.AnotacionGrupo).filter(
        models.AnotacionGrupo.grupo_id == grupo_id
    ).first()

    if not anotacion:
        anotacion = models.AnotacionGrupo(grupo_id=grupo_id)
        db.add(anotacion)

    for field, value in body.model_dump(exclude_unset=True).items():
        setattr(anotacion, field, value)

    anotacion.actualizado_en = datetime.now().isoformat()
    db.commit()
    db.refresh(anotacion)

    out = schemas.AnotacionGrupoOut.model_validate(anotacion)
    out.grupo_nombre = grupo.nombre
    out.empresa_nombre = grupo.empresa.nombre if grupo.empresa else None
    return out


# ─────────────────────────────────────────
# POST /anotaciones-grupo/{grupo_id}/hoteles
# Agrega segmento de hotel
# ─────────────────────────────────────────
@router.post("/{grupo_id}/hoteles", response_model=schemas.HotelSegmentoOut)
def add_hotel_segmento(grupo_id: int, body: schemas.HotelSegmentoCreate, db: Session = Depends(get_db)):
    anotacion = db.query(models.AnotacionGrupo).filter(
        models.AnotacionGrupo.grupo_id == grupo_id
    ).first()

    if not anotacion:
        grupo = db.query(models.Grupo).filter(models.Grupo.id == grupo_id).first()
        if not grupo:
            raise HTTPException(status_code=404, detail="Grupo no encontrado")
        anotacion = models.AnotacionGrupo(grupo_id=grupo_id)
        db.add(anotacion)
        db.commit()
        db.refresh(anotacion)

    segmento = models.HotelGrupoSegmento(
        anotacion_id=anotacion.id,
        **body.model_dump()
    )
    db.add(segmento)
    db.commit()
    db.refresh(segmento)
    return segmento


# ─────────────────────────────────────────
# PUT /anotaciones-grupo/hoteles/{segmento_id}
# Edita segmento de hotel
# ─────────────────────────────────────────
@router.put("/hoteles/{segmento_id}", response_model=schemas.HotelSegmentoOut)
def update_hotel_segmento(segmento_id: int, body: schemas.HotelSegmentoCreate, db: Session = Depends(get_db)):
    segmento = db.query(models.HotelGrupoSegmento).filter(
        models.HotelGrupoSegmento.id == segmento_id
    ).first()
    if not segmento:
        raise HTTPException(status_code=404, detail="Segmento no encontrado")

    for field, value in body.model_dump(exclude_unset=True).items():
        setattr(segmento, field, value)

    db.commit()
    db.refresh(segmento)
    return segmento


# ─────────────────────────────────────────
# DELETE /anotaciones-grupo/hoteles/{segmento_id}
# Elimina segmento de hotel
# ─────────────────────────────────────────
@router.delete("/hoteles/{segmento_id}")
def delete_hotel_segmento(segmento_id: int, db: Session = Depends(get_db)):
    segmento = db.query(models.HotelGrupoSegmento).filter(
        models.HotelGrupoSegmento.id == segmento_id
    ).first()
    if not segmento:
        raise HTTPException(status_code=404, detail="Segmento no encontrado")

    db.delete(segmento)
    db.commit()
    return {"ok": True}


# ─────────────────────────────────────────
# POST /anotaciones-grupo/{grupo_id}/comidas
# Agrega comida
# ─────────────────────────────────────────
@router.post("/{grupo_id}/comidas", response_model=schemas.ComidaGrupoOut)
def add_comida(grupo_id: int, body: schemas.ComidaGrupoCreate, db: Session = Depends(get_db)):
    anotacion = db.query(models.AnotacionGrupo).filter(
        models.AnotacionGrupo.grupo_id == grupo_id
    ).first()

    if not anotacion:
        grupo = db.query(models.Grupo).filter(models.Grupo.id == grupo_id).first()
        if not grupo:
            raise HTTPException(status_code=404, detail="Grupo no encontrado")
        anotacion = models.AnotacionGrupo(grupo_id=grupo_id)
        db.add(anotacion)
        db.commit()
        db.refresh(anotacion)

    comida = models.ComidaGrupo(
        anotacion_id=anotacion.id,
        **body.model_dump()
    )
    db.add(comida)
    db.commit()
    db.refresh(comida)
    return comida


# ─────────────────────────────────────────
# PUT /anotaciones-grupo/comidas/{comida_id}
# Edita comida
# ─────────────────────────────────────────
@router.put("/comidas/{comida_id}", response_model=schemas.ComidaGrupoOut)
def update_comida(comida_id: int, body: schemas.ComidaGrupoCreate, db: Session = Depends(get_db)):
    comida = db.query(models.ComidaGrupo).filter(
        models.ComidaGrupo.id == comida_id
    ).first()
    if not comida:
        raise HTTPException(status_code=404, detail="Comida no encontrada")

    for field, value in body.model_dump(exclude_unset=True).items():
        setattr(comida, field, value)

    db.commit()
    db.refresh(comida)
    return comida


# ─────────────────────────────────────────
# DELETE /anotaciones-grupo/comidas/{comida_id}
# Elimina comida
# ─────────────────────────────────────────
@router.delete("/comidas/{comida_id}")
def delete_comida(comida_id: int, db: Session = Depends(get_db)):
    comida = db.query(models.ComidaGrupo).filter(
        models.ComidaGrupo.id == comida_id
    ).first()
    if not comida:
        raise HTTPException(status_code=404, detail="Comida no encontrada")

    db.delete(comida)
    db.commit()
    return {"ok": True}


# ─────────────────────────────────────────
# GET /anotaciones-grupo/{grupo_id}/rooming-excel
# Genera y descarga Excel de rooming list profesional
# ─────────────────────────────────────────
@router.get("/{grupo_id}/rooming-excel")
def rooming_excel(grupo_id: int, db: Session = Depends(get_db)):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import (
            PatternFill, Font, Alignment, Border, Side
        )
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise HTTPException(
            status_code=500,
            detail="openpyxl no está instalado. Ejecute: pip install openpyxl"
        )

    grupo = db.query(models.Grupo).filter(models.Grupo.id == grupo_id).first()
    if not grupo:
        raise HTTPException(status_code=404, detail="Grupo no encontrado")

    anotacion = db.query(models.AnotacionGrupo).filter(
        models.AnotacionGrupo.grupo_id == grupo_id
    ).first()

    empresa_nombre = grupo.empresa.nombre if grupo.empresa else "—"
    grupo_nombre = grupo.nombre or "—"
    fecha_entrada = grupo.fecha_entrada.strftime("%d/%m/%Y") if grupo.fecha_entrada else "—"
    fecha_salida = grupo.fecha_salida.strftime("%d/%m/%Y") if grupo.fecha_salida else "—"

    hoteles = anotacion.hoteles if anotacion else []
    comidas = sorted(anotacion.comidas, key=lambda c: c.fecha) if anotacion else []
    cant_hombres = anotacion.cant_hombres if anotacion else None
    cant_mujeres = anotacion.cant_mujeres if anotacion else None
    cant_tripulacion = anotacion.cant_tripulacion if anotacion else None
    notas_generales = anotacion.notas_generales if anotacion else None

    total_pax = (cant_hombres or 0) + (cant_mujeres or 0) + (cant_tripulacion or 0)

    # ── Estilos ──────────────────────────────────────────────────
    NAVY = "1e3a5f"
    LIGHT_BLUE = "d6e4f0"
    LIGHT_GRAY = "f2f2f2"
    WHITE = "FFFFFF"

    header_fill = PatternFill("solid", fgColor=NAVY)
    header_font = Font(color=WHITE, bold=True, size=11)
    title_font = Font(color=NAVY, bold=True, size=16)
    subtitle_font = Font(color=NAVY, bold=True, size=12)
    section_font = Font(color=WHITE, bold=True, size=11)
    section_fill = PatternFill("solid", fgColor=NAVY)
    alt_fill = PatternFill("solid", fgColor=LIGHT_GRAY)
    normal_fill = PatternFill("solid", fgColor=WHITE)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    thin = Side(style="thin", color="AAAAAA")
    thick = Side(style="medium", color=NAVY)
    thin_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    thick_border = Border(left=thick, right=thick, top=thick, bottom=thick)

    def col_border(left_s=thin, right_s=thin, top_s=thin, bottom_s=thin):
        return Border(
            left=Side(style=left_s, color="AAAAAA"),
            right=Side(style=right_s, color="AAAAAA"),
            top=Side(style=top_s, color="AAAAAA"),
            bottom=Side(style=bottom_s, color="AAAAAA"),
        )

    # ── Workbook ─────────────────────────────────────────────────
    wb = Workbook()
    ws = wb.active
    ws.title = "Rooming List"

    # Definir anchos de columna (A–E)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 34

    # ── HEADER PRINCIPAL ─────────────────────────────────────────
    # Fila 1: Título
    ws.merge_cells("A1:E1")
    ws["A1"] = "ROOMING LIST"
    ws["A1"].font = Font(color=WHITE, bold=True, size=20)
    ws["A1"].fill = header_fill
    ws["A1"].alignment = center
    ws.row_dimensions[1].height = 36

    # Fila 2: Empresa + Grupo
    ws.merge_cells("A2:B2")
    ws["A2"] = f"Empresa: {empresa_nombre}"
    ws["A2"].font = Font(color=WHITE, bold=True, size=12)
    ws["A2"].fill = PatternFill("solid", fgColor="2d5a8e")
    ws["A2"].alignment = left

    ws.merge_cells("C2:E2")
    ws["C2"] = f"Grupo: {grupo_nombre}"
    ws["C2"].font = Font(color=WHITE, bold=True, size=12)
    ws["C2"].fill = PatternFill("solid", fgColor="2d5a8e")
    ws["C2"].alignment = left
    ws.row_dimensions[2].height = 24

    # Fila 3: Fechas + PAX total
    ws.merge_cells("A3:B3")
    ws["A3"] = f"Check-in: {fecha_entrada}   |   Check-out: {fecha_salida}"
    ws["A3"].font = Font(color=WHITE, size=11)
    ws["A3"].fill = PatternFill("solid", fgColor="3a6fa8")
    ws["A3"].alignment = left

    ws.merge_cells("C3:E3")
    ws["C3"] = f"Total PAX: {total_pax if total_pax else '—'}"
    ws["C3"].font = Font(color=WHITE, bold=True, size=11)
    ws["C3"].fill = PatternFill("solid", fgColor="3a6fa8")
    ws["C3"].alignment = left
    ws.row_dimensions[3].height = 22

    # Espacio
    ws.row_dimensions[4].height = 10

    # ── SECCIÓN: COMPOSICIÓN DEL GRUPO ───────────────────────────
    row = 5
    ws.merge_cells(f"A{row}:E{row}")
    ws[f"A{row}"] = "COMPOSICIÓN DEL GRUPO"
    ws[f"A{row}"].font = section_font
    ws[f"A{row}"].fill = section_fill
    ws[f"A{row}"].alignment = center
    ws.row_dimensions[row].height = 22
    row += 1

    composicion = [
        ("Hombres", cant_hombres if cant_hombres is not None else "—"),
        ("Mujeres", cant_mujeres if cant_mujeres is not None else "—"),
        ("Tripulación (guías/choferes)", cant_tripulacion if cant_tripulacion is not None else "—"),
        ("Total PAX", total_pax if total_pax else "—"),
    ]
    for i, (label, val) in enumerate(composicion):
        fill = alt_fill if i % 2 else normal_fill
        ws.merge_cells(f"A{row}:C{row}")
        ws[f"A{row}"] = label
        ws[f"A{row}"].fill = fill
        ws[f"A{row}"].border = thin_border
        ws[f"A{row}"].alignment = left
        ws[f"A{row}"].font = Font(bold=(label == "Total PAX"))

        ws.merge_cells(f"D{row}:E{row}")
        ws[f"D{row}"] = val
        ws[f"D{row}"].fill = fill
        ws[f"D{row}"].border = thin_border
        ws[f"D{row}"].alignment = center
        ws[f"D{row}"].font = Font(bold=(label == "Total PAX"))
        ws.row_dimensions[row].height = 20
        row += 1

    row += 1  # espacio

    # ── SECCIÓN: HOTELES ─────────────────────────────────────────
    ws.merge_cells(f"A{row}:E{row}")
    ws[f"A{row}"] = "HOTELES"
    ws[f"A{row}"].font = section_font
    ws[f"A{row}"].fill = section_fill
    ws[f"A{row}"].alignment = center
    ws.row_dimensions[row].height = 22
    row += 1

    # Encabezado tabla hoteles
    hotel_headers = ["Hotel", "PAX", "Check-in", "Check-out", "Notas"]
    for col_i, h in enumerate(hotel_headers, start=1):
        cell = ws.cell(row=row, column=col_i, value=h)
        cell.font = header_font
        cell.fill = PatternFill("solid", fgColor="2d5a8e")
        cell.alignment = center
        cell.border = thin_border
    ws.row_dimensions[row].height = 20
    row += 1

    if hoteles:
        for i, seg in enumerate(hoteles):
            fill = alt_fill if i % 2 else normal_fill
            vals = [
                seg.hotel_nombre,
                seg.cant_pax if seg.cant_pax is not None else "—",
                seg.fecha_ingreso.strftime("%d/%m/%Y") if seg.fecha_ingreso else "—",
                seg.fecha_salida.strftime("%d/%m/%Y") if seg.fecha_salida else "—",
                seg.notas or "",
            ]
            for col_i, v in enumerate(vals, start=1):
                cell = ws.cell(row=row, column=col_i, value=v)
                cell.fill = fill
                cell.border = thin_border
                cell.alignment = left if col_i in (1, 5) else center
            ws.row_dimensions[row].height = 20
            row += 1
    else:
        ws.merge_cells(f"A{row}:E{row}")
        ws[f"A{row}"] = "Sin segmentos de hotel registrados"
        ws[f"A{row}"].alignment = center
        ws[f"A{row}"].font = Font(italic=True, color="888888")
        ws.row_dimensions[row].height = 20
        row += 1

    row += 1  # espacio

    # ── SECCIÓN: PROGRAMA DE COMIDAS ─────────────────────────────
    ws.merge_cells(f"A{row}:E{row}")
    ws[f"A{row}"] = "PROGRAMA DE COMIDAS"
    ws[f"A{row}"].font = section_font
    ws[f"A{row}"].fill = section_fill
    ws[f"A{row}"].alignment = center
    ws.row_dimensions[row].height = 22
    row += 1

    comida_headers = ["Fecha", "Comida", "Lugar", "PAX", "Notas"]
    for col_i, h in enumerate(comida_headers, start=1):
        cell = ws.cell(row=row, column=col_i, value=h)
        cell.font = header_font
        cell.fill = PatternFill("solid", fgColor="2d5a8e")
        cell.alignment = center
        cell.border = thin_border
    ws.row_dimensions[row].height = 20
    row += 1

    if comidas:
        for i, comida in enumerate(comidas):
            fill = alt_fill if i % 2 else normal_fill
            vals = [
                comida.fecha.strftime("%d/%m/%Y"),
                comida.tipo_comida.capitalize(),
                comida.lugar,
                comida.cant_pax if comida.cant_pax is not None else "—",
                comida.notas or "",
            ]
            for col_i, v in enumerate(vals, start=1):
                cell = ws.cell(row=row, column=col_i, value=v)
                cell.fill = fill
                cell.border = thin_border
                cell.alignment = left if col_i in (3, 5) else center
            ws.row_dimensions[row].height = 20
            row += 1
    else:
        ws.merge_cells(f"A{row}:E{row}")
        ws[f"A{row}"] = "Sin comidas registradas"
        ws[f"A{row}"].alignment = center
        ws[f"A{row}"].font = Font(italic=True, color="888888")
        ws.row_dimensions[row].height = 20
        row += 1

    row += 1  # espacio

    # ── SECCIÓN: NOTAS GENERALES ──────────────────────────────────
    ws.merge_cells(f"A{row}:E{row}")
    ws[f"A{row}"] = "NOTAS GENERALES"
    ws[f"A{row}"].font = section_font
    ws[f"A{row}"].fill = section_fill
    ws[f"A{row}"].alignment = center
    ws.row_dimensions[row].height = 22
    row += 1

    ws.merge_cells(f"A{row}:E{row + 3}")
    ws[f"A{row}"] = notas_generales or ""
    ws[f"A{row}"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    ws[f"A{row}"].border = thin_border
    ws[f"A{row}"].fill = normal_fill
    ws.row_dimensions[row].height = 18
    ws.row_dimensions[row + 1].height = 18
    ws.row_dimensions[row + 2].height = 18
    ws.row_dimensions[row + 3].height = 18
    row += 5

    # ── FOOTER ───────────────────────────────────────────────────
    ws.merge_cells(f"A{row}:E{row}")
    ws[f"A{row}"] = "Generado por Sistema de Gestión de Eventos"
    ws[f"A{row}"].font = Font(italic=True, color="888888", size=9)
    ws[f"A{row}"].alignment = center
    ws.row_dimensions[row].height = 16

    # ── Serializar y devolver ─────────────────────────────────────
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    safe_nombre = grupo_nombre.replace(" ", "_")
    filename = f"rooming_{safe_nombre}.xlsx"

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
